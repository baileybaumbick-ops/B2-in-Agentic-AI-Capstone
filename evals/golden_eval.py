"""Golden-50 evaluation (code only, no model calls).

    python evals/golden_eval.py import                       # xlsx -> golden_50_human.csv + benchmark json
    python evals/golden_eval.py score --run-dir runs/full     # compare saved predictions vs human labels

Human labels are never sent to a model; this script only reads saved predictions.
Sentiment tolerance is predeclared: |pred - human| <= 0.5 counts as agreement; MAE is also reported.
"""
import argparse
import csv
import gzip
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from pipeline.common import INTENTS, TOPICS, read_csv, write_json  # noqa: E402

GOLD_DIR = ROOT / "evals" / "golden"
SHEET = GOLD_DIR / "golden_50_label_sheet.xlsx"
HUMAN_CSV = GOLD_DIR / "golden_50_human.csv"
BENCH = GOLD_DIR / "golden_50_benchmark.json"
SENTIMENT_TOLERANCE = 0.5


def _bool(v):
    return str(v).strip().upper() in ("TRUE", "1", "YES", "Y")


def cmd_import(a):
    from openpyxl import load_workbook
    ws = load_workbook(SHEET, data_only=True)["label_here"]
    header = [c.value for c in ws[1]]
    rows, problems = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, r)}
        if not d["review_id"]:
            continue
        for f, allowed in (("topic", TOPICS), ("intent", INTENTS)):
            if d[f] not in allowed:
                problems.append(f"{d['#']}: {f}={d[f]!r}")
        if d["severity"] not in "12345" or not d["severity"]:
            problems.append(f"{d['#']}: severity={d['severity']!r}")
        if d["evidence_quote"] and d["evidence_quote"] not in d["review_text"]:
            problems.append(f"{d['#']}: evidence_quote is not an exact substring")
        rows.append(d)
    if problems:
        print("Fix these cells in the sheet first:\n  " + "\n  ".join(problems))
        sys.exit(1)
    with HUMAN_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=header, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    cases = []
    for d in rows:
        def acc(field, cast=str):
            vals = [cast(d[field])] + ([cast(d[f"also_acceptable_{field}"])] if d.get(f"also_acceptable_{field}") else [])
            return list(dict.fromkeys(vals))
        cases.append({"review_id": d["review_id"], "status": "approved", "reviewer": a.reviewer,
                      "accepted_topic": acc("topic"), "accepted_intent": acc("intent"), "accepted_severity": acc("severity", int)})
    write_json(BENCH, {"version": "golden-50-human-v1", "cases": cases})
    print(f"Imported {len(rows)} human labels -> {HUMAN_CSV.name}, {BENCH.name}")


def load_predictions(run_dir):
    with gzip.open(Path(run_dir) / "enriched.jsonl.gz", "rt", encoding="utf-8") as f:
        return {r["review_id"]: r for r in map(json.loads, f)}


def cmd_score(a):
    human = list(csv.DictReader(HUMAN_CSV.open(encoding="utf-8", newline="")))
    preds = load_predictions(a.run_dir)
    texts = {r["review_id"]: r["review_text"] for r in read_csv(ROOT / "data" / "golden_50_to_label.csv")}
    tot = Counter()
    conf = {"topic": defaultdict(Counter), "intent": defaultdict(Counter)}
    per_topic = defaultdict(Counter)
    cases = []
    for h in human:
        rid = h["review_id"]
        p = preds.get(rid)
        ok_pred = bool(p and p.get("status") == "completed")
        acc_topic = {h["topic"]} | ({h["also_acceptable_topic"]} if h.get("also_acceptable_topic") else set())
        acc_intent = {h["intent"]} | ({h["also_acceptable_intent"]} if h.get("also_acceptable_intent") else set())
        acc_sev = {int(h["severity"])} | ({int(h["also_acceptable_severity"])} if h.get("also_acceptable_severity") else set())
        c = {"review_id": rid, "review_text": texts.get(rid, "")[:200], "human_topic": h["topic"], "human_subtopic": h.get("subtopic"),
             "human_intent": h["intent"], "human_severity": h["severity"], "human_sentiment": h["sentiment"],
             "ambiguous": _bool(h.get("ambiguous")), "prediction_status": p.get("status") if p else "missing"}
        if ok_pred:
            sev_err = min(abs(p["severity"] - s) for s in acc_sev)
            sent_err = abs(p["sentiment"] - float(h["sentiment"])) if h["sentiment"] else None
            ents = p.get("entities", [])
            unsupported = [e for e in ents if e.lower() not in texts.get(rid, "").lower()]
            c.update(pred_topic=p["topic"], pred_subtopic=p.get("subtopic"), pred_intent=p["intent"], pred_severity=p["severity"],
                     pred_sentiment=p["sentiment"], pred_quote=p["evidence_quote"], pred_needs_review=p["needs_review"],
                     topic_ok=p["topic"] in acc_topic, intent_ok=p["intent"] in acc_intent, severity_exact=p["severity"] in acc_sev,
                     severity_abs_err=sev_err, sentiment_abs_err=sent_err,
                     sentiment_ok=None if sent_err is None else sent_err <= SENTIMENT_TOLERANCE,
                     quote_exact_substring=p["evidence_quote"] in texts.get(rid, ""), entities=ents, unsupported_entities=unsupported,
                     quote_supports_label="")  # filled by human inspection in the CSV
            tot["predicted"] += 1
            for k in ("topic_ok", "intent_ok", "severity_exact", "quote_exact_substring"):
                tot[k] += c[k]
            tot["all3"] += c["topic_ok"] and c["intent_ok"] and c["severity_exact"]
            tot["sev_err"] += sev_err
            if sent_err is not None:
                tot["sent_n"] += 1
                tot["sent_err"] += sent_err
                tot["sentiment_ok"] += c["sentiment_ok"]
            tot["entities"] += len(ents)
            tot["unsupported_entities"] += len(unsupported)
            nr, amb = bool(p["needs_review"]), c["ambiguous"]
            tot["nr_tp"] += nr and amb
            tot["nr_fp"] += nr and not amb
            tot["nr_fn"] += (not nr) and amb
            conf["topic"][h["topic"]][p["topic"]] += 1
            conf["intent"][h["intent"]][p["intent"]] += 1
            per_topic[h["topic"]]["support"] += 1
            per_topic[h["topic"]]["correct"] += c["topic_ok"]
        else:
            for k in ("topic_ok", "intent_ok", "severity_exact"):
                c[k] = False  # missing/quarantined count as incorrect
        tot["ambiguous"] += c["ambiguous"]
        cases.append(c)
    n = len(human)

    def rate(k, d=n):
        return round(tot[k] / d, 4) if d else None
    results = {"run_dir": str(a.run_dir), "cases": n, "valid_predictions": tot["predicted"], "ambiguous_cases": tot["ambiguous"],
               "topic_agreement": rate("topic_ok"), "intent_agreement": rate("intent_ok"), "severity_exact_agreement": rate("severity_exact"),
               "all_three_agreement": rate("all3"), "severity_mae": rate("sev_err", tot["predicted"]),
               "sentiment_mae": rate("sent_err", tot["sent_n"]), "sentiment_within_tolerance": rate("sentiment_ok", tot["sent_n"]),
               "sentiment_tolerance": SENTIMENT_TOLERANCE, "quote_exact_substring_rate": rate("quote_exact_substring", tot["predicted"]),
               "entities_predicted": tot["entities"], "entities_not_found_in_text": tot["unsupported_entities"],
               "needs_review_vs_human_ambiguous": {"tp": tot["nr_tp"], "fp": tot["nr_fp"], "fn": tot["nr_fn"],
                                                   "precision": round(tot["nr_tp"] / (tot["nr_tp"] + tot["nr_fp"]), 4) if tot["nr_tp"] + tot["nr_fp"] else None,
                                                   "recall": round(tot["nr_tp"] / (tot["nr_tp"] + tot["nr_fn"]), 4) if tot["nr_tp"] + tot["nr_fn"] else None},
               "per_topic": {t: dict(v) for t, v in sorted(per_topic.items())},
               "confusion_topic_human_to_pred": {k: dict(v) for k, v in conf["topic"].items()},
               "confusion_intent_human_to_pred": {k: dict(v) for k, v in conf["intent"].items()},
               "note": "50 cases is a small diagnostic sample, not a precise population accuracy estimate."}
    out = ROOT / "evals" / "results"
    write_json(out / "golden_eval.json", results)
    keys = list(dict.fromkeys(k for c in cases for k in c))
    with (out / "golden_per_case.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=keys, lineterminator="\n")
        w.writeheader()
        w.writerows(cases)
    print(json.dumps({k: v for k, v in results.items() if not k.startswith("confusion")}, indent=2))


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    i = sub.add_parser("import")
    i.add_argument("--reviewer", default="student")
    i.set_defaults(fn=cmd_import)
    s = sub.add_parser("score")
    s.add_argument("--run-dir", required=True)
    s.set_defaults(fn=cmd_score)
    a = p.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
