"""Build the human labeling workbook for the golden 50 (no model involved).

    python evals/make_label_sheet.py          -> evals/golden/golden_50_label_sheet.xlsx

The human fills it in; evals/import_golden.py converts it to golden_50_human.csv/json.
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from pipeline import taxonomy as tx  # noqa: E402
from pipeline.common import INTENTS, TOPICS  # noqa: E402

SRC = ROOT / "data" / "golden_50_to_label.csv"
OUT = ROOT / "evals" / "golden" / "golden_50_label_sheet.xlsx"
COLS = ["#", "review_id", "review_text", "topic", "subtopic", "intent", "severity", "sentiment",
        "evidence_quote", "entities", "needs_review", "ambiguous", "also_acceptable_topic", "also_acceptable_intent",
        "also_acceptable_severity", "notes"]


def main():
    rows = list(csv.DictReader(SRC.open(encoding="utf-8", newline="")))
    wb = Workbook()
    ws = wb.active
    ws.title = "label_here"
    ws.append(COLS)
    for i, r in enumerate(rows, 1):
        ws.append([i, r["review_id"], r["review_text"]] + [""] * (len(COLS) - 3))
    head = PatternFill("solid", fgColor="1DB954")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
    widths = {"A": 4, "B": 12, "C": 70, "D": 12, "E": 22, "F": 13, "G": 9, "H": 10, "I": 40, "J": 20, "K": 13,
              "L": 11, "M": 14, "N": 14, "O": 14, "P": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_col=3):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    n = len(rows) + 1
    subs = [s for v in tx.SUBTOPICS.values() for s in v]
    lists = {"D": TOPICS, "M": TOPICS, "E": subs, "F": INTENTS, "N": INTENTS, "G": ["1", "2", "3", "4", "5"],
             "O": ["1", "2", "3", "4", "5"], "H": ["-1", "-0.5", "0", "0.5", "1"], "K": ["TRUE", "FALSE"], "L": ["TRUE", "FALSE"]}
    lists_ws = wb.create_sheet("lists")
    for j, (col, values) in enumerate(lists.items(), 1):
        for i, v in enumerate(values, 1):
            lists_ws.cell(row=i, column=j, value=v)
        letter = lists_ws.cell(row=1, column=j).column_letter
        dv = DataValidation(type="list", formula1=f"=lists!${letter}$1:${letter}${len(values)}", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    ws.freeze_panes = "D2"
    d = wb.create_sheet("definitions", 0)
    d.column_dimensions["A"].width = 120
    text = ["HOW TO LABEL (label the text only; ignore stars; you are the human reference)", "",
            "1. Pick topic + subtopic, intent, severity, sentiment for each review.",
            "2. evidence_quote: copy an EXACT piece of the review text that supports your label.",
            "3. entities: named features/devices/plans/competitors, separated by ';' (blank if none).",
            "4. needs_review / ambiguous: TRUE if reasonable people could label it differently;",
            "   then optionally fill also_acceptable_* with the other defensible label.",
            "5. Do NOT ask an AI to fill this in - these are the human labels the model is graded against.", ""]
    text += tx.definitions_text().splitlines()
    text += ["", "TIE RULES: highest supported severity wins; on a tie, the first specific problem mentioned.",
             "Positive review: first specific praised feature; general praise = other/general_praise.",
             "Paid-plan mention alone is not billing. Free-tier limits (skips, shuffle) = billing/free_tier_restrictions.",
             "Bare boycott slogans = unclear (topic other/boycott_or_offtopic) unless a product complaint or departure is stated."]
    for i, line in enumerate(text, 1):
        d.cell(row=i, column=1, value=line)
    d["A1"].font = Font(bold=True, size=13)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("Wrote", OUT, f"({len(rows)} reviews)")


if __name__ == "__main__":
    main()
